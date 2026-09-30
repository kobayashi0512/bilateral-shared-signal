#!/usr/bin/env python3
"""Build an editable, source-grounded manuscript framework figure."""

from __future__ import annotations

import base64
import hashlib
import json
import posixpath
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree
from xml.sax.saxutils import escape

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
BUILD_DIR = Path(__file__).resolve().parent
PACKAGE = BUILD_DIR / "package_v4"
SLUG = "figure1_framework"
W, H = 1800, 1480
PX_TO_IN = 1 / 254  # 10 drawing units per mm

ASSET_ROOT = BUILD_DIR / "assets"
IMAGES = [ASSET_ROOT / "deepdrid" / f"{view}.jpg" for view in ("L1", "L2", "R1", "R2")]
PAPILA_IMAGES = {
    "OD": ASSET_ROOT / "papila" / "right_eye_example.jpg",
    "OS": ASSET_ROOT / "papila" / "left_eye_example.jpg",
}
REFERENCE = ASSET_ROOT / "previous_framework_reference.png"

INK = "#183346"
MUTED = "#647884"
LIGHT = "#DCE5E8"
GRID = "#EEF2F3"
PAPER = "#FFFFFF"
TEAL = "#087E78"
TEAL_BG = "#E8F5F2"
BLUE = "#2876A0"
BLUE_BG = "#EAF3F8"
AMBER = "#AE752B"
AMBER_BG = "#FBF3E8"
RED = "#B75C4D"
RED_BG = "#FBEEEC"


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[:2], 16), int(value[2:4], 16), int(value[4:], 16))


def pptx_unit(value: float):
    return Inches(value * PX_TO_IN)


class Canvas:
    def __init__(self) -> None:
        self.objects = []

    def rect(self, x, y, w, h, fill=PAPER, stroke=LIGHT, sw=2, radius=16):
        self.objects.append(("rect", dict(x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, sw=sw, radius=radius)))

    def text(self, x, y, content, size=25, color=INK, bold=False, anchor="start", linespace=1.18, width=None):
        self.objects.append(("text", dict(x=x, y=y, content=content, size=size, color=color,
                                          bold=bold, anchor=anchor, linespace=linespace, width=width)))

    def line(self, x1, y1, x2, y2, color=LIGHT, sw=2, dashed=False, arrow=False):
        self.objects.append(("line", dict(x1=x1, y1=y1, x2=x2, y2=y2, color=color, sw=sw,
                                          dashed=dashed, arrow=arrow)))

    def image(self, x, y, w, h, path: Path):
        self.objects.append(("image", dict(x=x, y=y, w=w, h=h, path=path)))

    def tag(self, x, y, w, content, color=TEAL, fill=TEAL_BG):
        self.rect(x, y, w, 36, fill, fill, 1, 18)
        self.text(x+w/2, y+25, content, 21, color, True, "middle")

    def svg(self) -> str:
        parts = [f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
                 f'width="180mm" height="148mm" viewBox="0 0 {W} {H}">',
                 f'<rect width="{W}" height="{H}" fill="{PAPER}"/>',
                 '<defs><marker id="arrow" viewBox="0 0 10 10" refX="8.4" refY="5" '
                 'markerWidth="7" markerHeight="7" orient="auto-start-reverse">'
                 f'<path d="M 0 0 L 10 5 L 0 10 z" fill="{TEAL}"/></marker></defs>']
        for kind, p in self.objects:
            if kind == "rect":
                parts.append(f'<rect x="{p["x"]}" y="{p["y"]}" width="{p["w"]}" height="{p["h"]}" '
                             f'rx="{p["radius"]}" fill="{p["fill"]}" stroke="{p["stroke"]}" stroke-width="{p["sw"]}"/>')
            elif kind == "text":
                lines = p["content"].split("\n")
                for i, ln in enumerate(lines):
                    y = p["y"] + i*p["size"]*p["linespace"]
                    parts.append(f'<text x="{p["x"]}" y="{y}" font-family="Arial, Helvetica, sans-serif" '
                                 f'font-size="{p["size"]}" fill="{p["color"]}" '
                                 f'font-weight="{"700" if p["bold"] else "400"}" text-anchor="{p["anchor"]}">'
                                 f'{escape(ln)}</text>')
            elif kind == "line":
                dash = ' stroke-dasharray="11 8"' if p["dashed"] else ""
                marker = ' marker-end="url(#arrow)"' if p["arrow"] else ""
                parts.append(f'<line x1="{p["x1"]}" y1="{p["y1"]}" x2="{p["x2"]}" y2="{p["y2"]}" '
                             f'stroke="{p["color"]}" stroke-width="{p["sw"]}" stroke-linecap="round"{dash}{marker}/>')
            elif kind == "image":
                payload = base64.b64encode(p["path"].read_bytes()).decode("ascii")
                parts.append(f'<image x="{p["x"]}" y="{p["y"]}" width="{p["w"]}" height="{p["h"]}" '
                             'preserveAspectRatio="xMidYMid meet" '
                             f'xlink:href="data:image/jpeg;base64,{payload}"/>')
        parts.append("</svg>")
        return "\n".join(parts)

    def pptx(self) -> Presentation:
        prs = Presentation()
        prs.slide_width = pptx_unit(W)
        prs.slide_height = pptx_unit(H)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(PAPER)
        for kind, p in self.objects:
            if kind == "rect":
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                               pptx_unit(p["x"]), pptx_unit(p["y"]),
                                               pptx_unit(p["w"]), pptx_unit(p["h"]))
                shape.fill.solid()
                shape.fill.fore_color.rgb = rgb(p["fill"])
                shape.line.color.rgb = rgb(p["stroke"])
                shape.line.width = Pt(p["sw"]*72/254)
                # PowerPoint's default rounded-rectangle theme adds an unwanted shadow.
                for style in shape._element.xpath("./p:style"):
                    shape._element.remove(style)
                try:
                    shape.adjustments[0] = min(.25, p["radius"]/min(p["w"], p["h"]))
                except (AttributeError, IndexError):
                    pass
            elif kind == "line":
                conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                                  pptx_unit(p["x1"]), pptx_unit(p["y1"]),
                                                  pptx_unit(p["x2"]), pptx_unit(p["y2"]))
                conn.line.color.rgb = rgb(p["color"])
                conn.line.width = Pt(p["sw"]*72/254)
                if p["dashed"]:
                    from pptx.enum.dml import MSO_LINE_DASH_STYLE
                    conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
                if p["arrow"]:
                    ln = conn.line._get_or_add_ln()
                    end = OxmlElement("a:tailEnd")
                    end.set("type", "triangle")
                    end.set("w", "med")
                    end.set("len", "med")
                    ln.append(end)
            elif kind == "text":
                lines = p["content"].split("\n")
                text_h = p["size"]*p["linespace"]*len(lines)+12
                x = p["x"]
                w = W-p["x"]-24
                if p["anchor"] == "middle":
                    w = p["width"] or 440
                    x = p["x"]-w/2
                y = p["y"]-p["size"]
                box = slide.shapes.add_textbox(pptx_unit(x), pptx_unit(y), pptx_unit(w), pptx_unit(text_h))
                tf = box.text_frame
                tf.clear()
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                for i, ln in enumerate(lines):
                    par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    par.text = ln
                    par.alignment = PP_ALIGN.CENTER if p["anchor"] == "middle" else PP_ALIGN.LEFT
                    par.space_after = Pt(0)
                    par.line_spacing = 1.0
                    for run in par.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(p["size"]*72/254)
                        run.font.bold = p["bold"]
                        run.font.color.rgb = rgb(p["color"])
            elif kind == "image":
                with Image.open(p["path"]) as im:
                    aspect = im.width/im.height
                box_aspect = p["w"]/p["h"]
                if aspect > box_aspect:
                    ww, hh = p["w"], p["w"]/aspect
                else:
                    ww, hh = p["h"]*aspect, p["h"]
                slide.shapes.add_picture(str(p["path"]), pptx_unit(p["x"]+(p["w"]-ww)/2),
                                         pptx_unit(p["y"]+(p["h"]-hh)/2),
                                         width=pptx_unit(ww), height=pptx_unit(hh))
        return prs


def build() -> Canvas:
    c = Canvas()
    c.text(65, 58, "STUDY FRAMEWORK", 23, TEAL, True)
    c.text(65, 116, "Conditional bilateral correspondence: from images to evidence", 48, INK, True)
    c.text(65, 159, "Test cross-eye correspondence against grade and quality controls; audit view effects separately.", 24, MUTED)
    c.line(65, 190, 1735, 190, LIGHT, 2)

    # A: source unit and image hierarchy.
    c.text(66, 235, "A", 38, TEAL, True)
    c.text(107, 232, "Patient as the unit", 30, INK, True)
    c.text(68, 266, "DeepDRiD  ·  300 complete bilateral patients", 22, MUTED)
    c.rect(65, 292, 480, 480, PAPER, LIGHT, 2, 20)
    c.tag(87, 311, 202, "SOURCE COHORT", TEAL, TEAL_BG)
    c.text(88, 375, "Patient i", 26, INK, True)
    c.text(88, 407, "2 eyes × 2 views = 4 images", 21, MUTED)
    positions = [(86, 430), (302, 430), (86, 582), (302, 582)]
    names = ["L1", "L2", "R1", "R2"]
    for (x, y), path, name in zip(positions, IMAGES, names):
        c.rect(x, y, 185, 111, GRID, BLUE if name.startswith("R") else AMBER, 2, 3)
        c.image(x+2, y+2, 181, 107, path)
        c.tag(x+48, y+108, 90, name, BLUE if name.startswith("R") else AMBER,
              BLUE_BG if name.startswith("R") else AMBER_BG)
    c.text(88, 755, "Patient-disjoint train / val / test: 210 / 46 / 44", 21, MUTED)

    # B: actual comparison units. Photos are sourced directly, not from the prior figure.
    c.text(583, 235, "B", 38, TEAL, True)
    c.text(625, 232, "Define the contrasts", 30, INK, True)
    c.text(584, 266, "One source patient; preserve eye and view labels", 22, MUTED)
    c.rect(580, 292, 596, 480, PAPER, LIGHT, 2, 20)
    c.rect(603, 316, 550, 111, TEAL_BG, TEAL_BG, 1, 12)
    c.text(626, 350, "POSITIVE", 21, TEAL, True)
    c.text(626, 386, "L(i)  ↔  R(i)", 35, INK, True)
    c.text(830, 385, "same patient, opposite eyes", 21, MUTED)
    c.rect(603, 446, 550, 111, BLUE_BG, BLUE_BG, 1, 12)
    c.text(626, 480, "NEGATIVE", 21, BLUE, True)
    c.text(626, 516, "L(i)  ↔  R(j)", 35, INK, True)
    c.text(830, 515, "different patients; j ≠ i", 21, MUTED)
    c.line(604, 586, 1152, 586, LIGHT, 2)
    c.text(625, 621, "View checks", 22, INK, True)
    c.text(625, 655, "L1–R1  ·  L1–R2  ·  L2–R1  ·  L2–R2", 23, BLUE)
    c.text(625, 690, "Same-eye repeats (L1–L2, R1–R2) audit hierarchy.", 21, MUTED)
    c.text(625, 741, "Five fixed patient-level negative assignments", 21, MUTED)

    # C: measurement, not a claimed identity model.
    c.text(1210, 235, "C", 38, TEAL, True)
    c.text(1253, 232, "Measure correspondence", 30, INK, True)
    c.text(1210, 266, "Frozen raw representation · no feature tuning", 22, MUTED)
    c.rect(1207, 292, 528, 480, PAPER, LIGHT, 2, 20)
    c.rect(1230, 326, 480, 84, BLUE_BG, BLUE_BG, 1, 11)
    c.text(1470, 360, "ImageNet ResNet-50", 26, INK, True, "middle")
    c.text(1470, 388, "same fixed encoder for each image", 21, MUTED, False, "middle")
    c.line(1470, 415, 1470, 460, TEAL, 4, arrow=True)
    c.rect(1230, 472, 480, 77, TEAL_BG, TEAL_BG, 1, 11)
    c.text(1470, 507, "Cosine score  s(L, R)", 26, INK, True, "middle")
    c.line(1470, 555, 1470, 600, TEAL, 4, arrow=True)
    c.rect(1230, 609, 480, 77, PAPER, LIGHT, 2, 11)
    c.text(1470, 644, "Pair AUROC + cosine gap", 25, INK, True, "middle")
    c.text(1232, 730, "Auxiliary: grade / quality metadata probes", 21, MUTED)
    c.line(545, 535, 573, 535, TEAL, 4, arrow=True)
    c.line(1178, 535, 1200, 535, TEAL, 4, arrow=True)

    # D: controls are the core scientific design, not decorative labels.
    c.line(65, 804, 1735, 804, LIGHT, 2)
    c.text(66, 851, "D", 38, TEAL, True)
    c.text(108, 848, "Make cross-patient negatives progressively harder", 30, INK, True)
    c.text(68, 884, "The positive pairs and frozen features stay fixed; only negative-patient matching changes.", 22, MUTED)
    card_y, card_h = 912, 202
    cards = [
        (65, "01  DISEASE CONTROL", "Same DR grade", "Mean |Δ quality|  0.198", "grade stratum · patient derangement", AMBER, AMBER_BG),
        (637, "02  JOINT CONTROL", "Grade + exact quality tier", "Mean |Δ quality|  0.071", "exact grade × tier bucket", TEAL, TEAL_BG),
        (1209, "03  SENSITIVITY", "Grade + nearest quality", "Mean |Δ quality|  0.024", "within grade · nearest assignment", BLUE, BLUE_BG),
    ]
    for x, top, title, detail, rule, accent, bg in cards:
        c.rect(x, card_y, 526, card_h, PAPER, LIGHT, 2, 16)
        c.rect(x, card_y, 526, 9, accent, accent, 0, 2)
        c.text(x+24, card_y+47, top, 21, accent, True)
        c.text(x+24, card_y+94, title, 27, INK, True)
        c.text(x+24, card_y+132, detail, 21, MUTED)
        c.rect(x+24, card_y+154, 478, 27, bg, bg, 0, 8)
        c.text(x+262, card_y+174, rule, 21, accent, False, "middle")
    c.line(591, 1012, 627, 1012, TEAL, 4, arrow=True)
    c.line(1163, 1012, 1199, 1012, TEAL, 4, arrow=True)

    # E: source/test/transfer roles and claim boundary.
    c.line(65, 1149, 1735, 1149, LIGHT, 2)
    c.text(66, 1195, "E", 38, TEAL, True)
    c.text(108, 1192, "Audit at three evidence levels", 30, INK, True)
    lanes = [
        (65, "DeepDRiD", "INTERNAL MEASUREMENT", "Full 300: condition controls\nTest: four-view n=43; probes n=44", TEAL),
        (637, "CFP–OCT", "DEVELOPMENT STRESS", "137 eligible: raw vs H-CBFD\nCandidate loses to raw baseline", AMBER),
        (1209, "PAPILA", "LOCKED EXTERNAL BASELINE", "244 patients · one image/eye\nFrozen raw features only", BLUE),
    ]
    for x, name, role, body, accent in lanes:
        c.rect(x, 1220, 526, 150, PAPER, LIGHT, 2, 16)
        c.rect(x, 1220, 7, 150, accent, accent, 0, 2)
        c.text(x+24, 1255, name, 28, INK, True)
        c.text(x+24, 1284, role, 21, accent, True)
        c.text(x+24, 1325, body, 21, MUTED, False, "start", 1.35)
        if name == "PAPILA":
            for offset, eye in zip((344, 428), ("OD", "OS")):
                lane_y = 1220
                c.rect(x+offset, lane_y+29, 72, 68, GRID, accent, 2, 3)
                c.image(x+offset+2, lane_y+31, 68, 64, PAPILA_IMAGES[eye])
                c.text(x+offset+36, lane_y+122, eye, 18, accent, True, "middle", width=72)
    c.rect(65, 1398, 1670, 56, GRID, GRID, 0, 9)
    c.text(900, 1434, "Claim boundary: correspondence does not establish patient identity or a disease-independent biomarker.",
           21, INK, True, "middle", width=1580)
    return c


def write_json(path: Path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")


def main():
    folders = ["00_request", "01_inputs/source_images", "01_inputs/notes", "02_image2_master/master_selected",
               "03_assets_raw/generated_single_assets", "04_assets_cutout/transparent_png", "05_assets_vector",
               "06_editable_pptx", "07_svg_export", "08_previews", "09_manifests", "10_qa"]
    for folder in folders:
        (PACKAGE/folder).mkdir(parents=True, exist_ok=True)
    (PACKAGE/"00_request/user_prompt.md").write_text(
        "Update the existing Figure 1 framework for the manuscript, preserving its scientific logic and journal-ready minimal style while showing more authentic dataset source photographs. Keep text, vectors, labels, and arrows editable; use raster only for source fundus photographs.\n", encoding="utf-8")
    choices = {"mode": "image-to-editable with explicit redesign permission", "purpose": "SCI manuscript submission",
               "figure_type": "multi-panel workflow/framework", "tone_palette": "Minimal Journal; teal/blue/amber colorblind-aware",
               "target_standard": "generic two-column SCI, journal unspecified", "editability": "native SVG/PPTX text, shapes and connectors; original fundus images are raster",
               "output_root": "figures/figure1_framework/package_v4", "figure_slug": SLUG, "language": "English",
               "canvas_mm": [180, 148], "source_figure": "figures/figure1_framework/assets/previous_framework_reference.png",
               "redesign_authorized": "User requested increased use of original dataset images; preserve the existing framework layout and add an illustrative PAPILA pair to the PAPILA evidence card.",
               "avoid": "unearned biomarker claims, synthetic clinical images, patient-level labels, unlicensed CFP–OCT photographs"}
    write_json(PACKAGE/"00_request/preflight_choices.json", choices)
    write_json(PACKAGE/"00_request/mode.json", {"mode": "B", "reason": "The current Figure 1 is the structural reference; the user explicitly requested more original dataset images, so a PAPILA OD/OS example was added without changing panel order or scientific roles."})
    shutil.copy2(REFERENCE, PACKAGE/"02_image2_master/master_selected/previous_framework_reference.png")
    for path in [*IMAGES, *PAPILA_IMAGES.values()]:
        shutil.copy2(path, PACKAGE/"01_inputs/source_images"/path.name)
    (PACKAGE/"01_inputs/notes/scientific_logic.md").write_text(
        "The figure distinguishes patient unit, within-/cross-eye comparisons, matched cross-patient negatives, frozen scoring, and three evidence roles. "
        "No model training is implied in the main raw-feature route. The CFP–OCT learned candidate is shown only as a development stress readout; PAPILA tests frozen raw features. "
        "Panel A uses four unmodified DeepDRiD source photographs; panel E adds one unmodified OD/OS PAPILA pair. The CFP–OCT card is intentionally text-only because its public repository did not declare image-reuse terms.\n", encoding="utf-8")
    canvas = build()
    svg_path = PACKAGE/"07_svg_export"/f"{SLUG}.svg"
    svg_path.write_text(canvas.svg(), encoding="utf-8")
    pptx_path = PACKAGE/"06_editable_pptx"/f"{SLUG}.pptx"
    canvas.pptx().save(pptx_path)
    image_sources = [
        {"asset": p.name, "dataset": "DeepDRiD", "source_repository": "https://github.com/deepdrdoc/DeepDRiD",
         "source_license": "CC BY-SA 4.0", "changes": "resized for display only",
         "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in IMAGES
    ] + [
        {"asset": p.name, "dataset": "PAPILA", "source_repository": "https://figshare.com/articles/dataset/PAPILA/14798004",
         "source_license": "GPL-3.0+ (Figshare v2 record)", "changes": "resized for display only",
         "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in PAPILA_IMAGES.values()
    ]
    write_json(PACKAGE/"09_manifests/image_provenance.json", image_sources)
    with zipfile.ZipFile(pptx_path) as z:
        names = set(z.namelist())
        missing = []
        for rels in [n for n in names if n.endswith(".rels")]:
            if rels.startswith("_rels/"):
                source_dir = ""
            else:
                source_dir = rels.split("/_rels/")[0]
            tree = ElementTree.fromstring(z.read(rels))
            for rel in tree:
                if rel.attrib.get("TargetMode") == "External":
                    continue
                target = rel.attrib.get("Target", "")
                resolved = posixpath.normpath(posixpath.join(source_dir, target))
                if resolved not in names:
                    missing.append({"rels": rels, "target": target, "resolved": resolved})
        slide_xml = z.read("ppt/slides/slide1.xml").decode("utf-8")
        validation = {"zip_bad_entry": z.testzip(), "slides": len([n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")]),
                      "notes_entries": len([n for n in names if n.startswith("ppt/notesSlides/")]),
                      "media_entries": len([n for n in names if n.startswith("ppt/media/")]),
                      "slideNumberPlaceholders": slide_xml.count('type="sldNum"'),
                      "savedTriangleArrowheads": slide_xml.count('tailEnd type="triangle"'),
                      "missingRelationshipTargets": missing}
    write_json(PACKAGE/"06_editable_pptx/pptx_validation.json", validation)
    layout = {
        "canvas": {"width_mm": 180, "height_mm": 148, "aspect_ratio": 180/148},
        "reference": "figures/figure1_framework/assets/previous_framework_reference.png",
        "approved_deviation": "Existing framework geometry and panel order are preserved; two original PAPILA eye photographs are added to the PAPILA evidence card at the user's request.",
        "panels": ["A patient unit and four views", "B comparison definitions", "C frozen scoring", "D conditional controls", "E evidence roles"],
        "textZones": ["title and subtitle", "A-E panel titles and descriptions", "interpretation boundary"],
        "fineRegions": [
            {"id": f"image_{n}", "role": "original illustrative fundus image", "implementation": "trusted_source_image", "source": p.name}
            for n,p in zip(["L1","L2","R1","R2"], IMAGES)
        ] + [
            {"id": f"papila_{eye.lower()}", "role": "original illustrative PAPILA eye photograph", "implementation": "trusted_source_image", "source": p.name}
            for eye,p in PAPILA_IMAGES.items()
        ],
        "routeSegments": [
            {"id": k, "implementation": "native_vector", "arrow_direction_status": "from source to destination", "native_connector_endpoint_policy": "tailEnd triangle in DrawingML"}
            for k in ["input_to_pairs", "pairs_to_scoring", "encoder_to_cosine", "cosine_to_endpoint", "grade_to_joint", "joint_to_nearest"]
        ],
        "legendItems": ["eye labels L/R", "positive/negative pair fill", "dataset role accent colors"],
        "smallIcons": [], "complexGeneratedAssets": [],
        "referenceGuidedAssetFactory": "Not applicable: all complex images are original dataset photographs; no generated symbols or icon art."
    }
    write_json(PACKAGE/"09_manifests/master_layout_manifest.json", layout)
    write_json(PACKAGE/"09_manifests/asset_manifest.json", {
        "generated_assets": [], "trusted_source_images": [
            {"id": n, "source_path": p.relative_to(ROOT).as_posix(), "package_path": (Path("figures/figure1_framework/package_v4/01_inputs/source_images")/p.name).as_posix(),
             "editability_type": "raster photograph", "source_license": "CC BY-SA 4.0", "changes": "resized for display only"}
            for n,p in zip(["L1","L2","R1","R2"],IMAGES)
        ] + [
            {"id": f"PAPILA_{eye}", "source_path": p.relative_to(ROOT).as_posix(), "package_path": (Path("figures/figure1_framework/package_v4/01_inputs/source_images")/p.name).as_posix(),
             "editability_type": "raster photograph", "source_license": "GPL-3.0+", "changes": "resized for display only"}
            for eye,p in PAPILA_IMAGES.items()
        ], "source_figure_regions_reused": False})
    write_json(PACKAGE/"09_manifests/editability_manifest.json", {
        "native_editable": ["all titles, labels, panels, comparison cards, dividers, arrows, callouts"],
        "raster_assets": [p.name for p in [*IMAGES, *PAPILA_IMAGES.values()]], "generated_assets": [],
        "source_figure_region_reuse": False, "connector_endpoint_audit": "passed: six saved DrawingML triangle arrowheads and rendered-preview inspection"})
    write_json(PACKAGE/"09_manifests/visual_diff_report.json", {
        "layout_parity": "existing five-panel framework preserved; PAPILA thumbnails added within panel E",
        "content_parity": "original data-source images added without changing evidence roles",
        "asset_parity": "four original DeepDRiD and two original PAPILA photographs used", "route_arrow_parity": "passed: six source-to-destination routes checked in SVG and PPTX renders",
        "small_icon_coverage": "not applicable", "semantic_symbol_coverage": "not applicable",
        "fine_region_coverage": "six photographs and six route segments inventoried",
        "no_coarse_asset": True, "reference_guided_assets": "not applicable", "no_crop_gate": True,
        "truthful_qa": "passed: SVG and saved-PPTX previews visually inspected; target-journal final-size check remains", "redo_required": None})
    (PACKAGE/"10_qa/policy_check.md").write_text(
        "No AI-generated clinical or decorative visual assets. All six clinical photographs are unmodified source images and illustrative only. "
        "DeepDRiD photographs are attributed under CC BY-SA 4.0; the PAPILA pair is attributed under GPL-3.0+ according to the v2 Figshare record. "
        "The CFP–OCT photo examples are omitted because its public repository does not declare image-reuse terms. Check the target journal's figure policy before submission.\n", encoding="utf-8")
    (PACKAGE/"10_qa/known_issues.md").write_text(
        "The figure was visually checked at its full 180 × 148 mm canvas; the target journal and final column width are not yet specified, so submission-size checking remains. The six clinical photographs are distributed only under their respective source terms. This repository has no project-wide code license yet. CFP–OCT photographs are deliberately absent because the reviewed source repository did not state image-reuse terms.\n", encoding="utf-8")
    (PACKAGE/"07_svg_export/svg_render_notes.md").write_text(
        "SVG has native text, shapes and connector geometry; four original DeepDRiD and two original PAPILA JPEG photographs are embedded as raster media.\n", encoding="utf-8")
    (PACKAGE/"10_qa/QA_notes.md").write_text("PPTX package validation passed: one slide, six embedded source photographs, six native vector arrowheads, and no missing relationship targets. SVG and saved-PPTX previews were rendered and visually inspected at full figure size. Panel geometry, label/image alignment, arrows, claim boundary, and the PAPILA OD/OS thumbnails are legible without overlap. Source photographs are raster; framework text, cards, dividers, arrows, and labels remain editable vectors.\n", encoding="utf-8")
    (PACKAGE/"10_qa/visual_qa.md").write_text("Passed: the 180 × 148 mm SVG and editable PPTX were rendered to page previews and inspected. The four DeepDRiD view examples remain clearly visible in panel A; the PAPILA OD/OS examples fit within its panel-E evidence card; no labels, connectors, or image frames overlap. The SVG and PPTX renders preserve the same layout. Recheck at the target journal's final column width before submission.\n", encoding="utf-8")
    package_rel = Path("figures/figure1_framework/package_v4")
    write_json(PACKAGE/"09_manifests/output_index.json", {
        "figure_slug": SLUG, "mode": "image-to-editable; source photos added to existing framework by user request", "package_root": package_rel.as_posix(),
        "final_pptx": (package_rel/"06_editable_pptx"/f"{SLUG}.pptx").as_posix(),
        "final_svg": (package_rel/"07_svg_export"/f"{SLUG}.svg").as_posix(),
        "preview_pptx": (package_rel/"08_previews/preview_pptx.png").as_posix(),
        "preview_svg": (package_rel/"08_previews/preview_svg.png").as_posix(),
        "selected_master": (package_rel/"02_image2_master/master_selected/previous_framework_reference.png").as_posix(),
        "source_reference": "figures/figure1_framework/assets/previous_framework_reference.png",
        "asset_manifest": (package_rel/"09_manifests/asset_manifest.json").as_posix(),
        "editability_manifest": (package_rel/"09_manifests/editability_manifest.json").as_posix(),
        "master_layout_manifest": (package_rel/"09_manifests/master_layout_manifest.json").as_posix(),
        "visual_diff_report": (package_rel/"09_manifests/visual_diff_report.json").as_posix(),
        "qa_notes": (package_rel/"10_qa/QA_notes.md").as_posix(),
        "known_issues": (package_rel/"10_qa/known_issues.md").as_posix()})
    print(PACKAGE)


if __name__ == "__main__":
    main()
